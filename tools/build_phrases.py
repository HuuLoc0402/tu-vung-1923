# Sinh phrases.js (100 cụm từ giao tiếp) từ file Excel trong vault:
#   tai-lieu/100 cụm từ tiếng Anh thông dụng.xlsx, sheet "Bản học"
# Cột B: cụm từ + xuống dòng + phiên âm; cột C: nghĩa; cột D: câu ví dụ + xuống dòng + câu dịch.
# Chạy: python tools/build_phrases.py "<đường dẫn file xlsx>"
import json, sys
from pathlib import Path
import openpyxl

src = Path(sys.argv[1])
ws = openpyxl.load_workbook(src)['Bản học']
rows = []
for r in range(3, ws.max_row + 1):
    b, c, d = ws[f'B{r}'].value, ws[f'C{r}'].value, ws[f'D{r}'].value
    if not b:
        continue
    en, ipa = b.split('\n', 1)
    ex, ex_vi = d.split('\n', 1)
    rows.append([en.strip(), ipa.strip(), c.strip(), ex.strip(), ex_vi.strip()])
assert len(rows) == 100, len(rows)
assert len({r[0].lower().rstrip('?!.') for r in rows}) == 100, 'trùng cụm từ'
out = Path(__file__).resolve().parent.parent / 'phrases.js'
body = ',\n'.join(json.dumps(r, ensure_ascii=False) for r in rows)
out.write_text(
    f'/* 100 cụm từ giao tiếp, sinh tự động từ tai-lieu/{src.name} trong vault (tools/build_phrases.py). [en, ipa, vi, ví dụ, dịch ví dụ] */\n'
    f'const PHRASES = [{body}];\n', encoding='utf-8')
print('saved', len(rows), out)
